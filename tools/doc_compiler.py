import os
from typing import Dict, Any, Optional
from pptx import Presentation
from pptx.util import Inches, Pt
from docx import Document
from docx.shared import Inches as DocxInches, Pt as DocxPt

class DocumentCompilerTool:
    """Tool that converts HLD Architecture JSON, Implementation Guidance, and rendered diagram images into PowerPoint (.pptx) or Word (.docx) documents."""
    
    def __init__(self, output_dir: str = "./tmp/deliverables"):
        self.output_dir = output_dir
        os.makedirs(self.output_dir, exist_ok=True)

    # =====================================================================
    # PowerPoint (.pptx) Compiler
    # =====================================================================
    def create_pptx_deck(
        self, 
        hld_data: Dict[str, Any], 
        diagram_paths: Optional[Dict[str, Optional[str]]] = None,
        implementation_guidance: Optional[Dict[str, Any]] = None,
        filename: str = "HLD_Presentation.pptx"
    ) -> str:
        """Generates a structured 16:9 PowerPoint presentation from HLD schema data, implementation specs, and diagrams."""
        diagrams = diagram_paths or {}
        prs = Presentation()
        
        # Widescreen 16:9 layout
        prs.slide_width = Inches(13.33)
        prs.slide_height = Inches(7.5)
        blank_layout = prs.slide_layouts[6]

        # Slide 1: Title Slide
        slide1 = prs.slides.add_slide(blank_layout)
        tb = slide1.shapes.add_textbox(Inches(1.5), Inches(2.5), Inches(10.33), Inches(2.0))
        tf = tb.text_frame
        p = tf.paragraphs[0]
        p.text = hld_data.get("system_name", "High-Level Architecture Design")
        p.font.bold = True
        p.font.size = Pt(40)
        
        p2 = tf.add_paragraph()
        p2.text = f"Pattern: {hld_data.get('architectural_pattern', 'N/A')}"
        p2.font.size = Pt(22)

        # Slide 2: Executive Summary & Core Components
        slide2 = prs.slides.add_slide(blank_layout)
        header_box = slide2.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(11.0), Inches(1.0))
        hp = header_box.text_frame.paragraphs[0]
        hp.text = "Executive Summary & System Components"
        hp.font.bold = True
        hp.font.size = Pt(28)

        content_box = slide2.shapes.add_textbox(Inches(0.8), Inches(1.5), Inches(11.5), Inches(5.0))
        ctf = content_box.text_frame
        ctf.word_wrap = True
        sp = ctf.paragraphs[0]
        sp.text = f"Summary: {hld_data.get('summary', '')}"
        sp.font.size = Pt(15)

        tf_comp_title = ctf.add_paragraph()
        tf_comp_title.text = "\nCore Components:"
        tf_comp_title.font.bold = True
        tf_comp_title.font.size = Pt(17)

        components = hld_data.get("components", [])
        for comp in components[:5]:
            cp = ctf.add_paragraph()
            cp.text = f"  {comp.get('name')} ({comp.get('tech_stack')}): {comp.get('description')}"
            cp.font.size = Pt(13)

        # Slide 3: Implementation Architecture Guidance (if available)
        if implementation_guidance:
            guidance_slide = prs.slides.add_slide(blank_layout)
            
            g_hdr = guidance_slide.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(11.0), Inches(0.8))
            gh_p = g_hdr.text_frame.paragraphs[0]
            gh_p.text = "Implementation Architecture & Tech Stack Matrix"
            gh_p.font.bold = True
            gh_p.font.size = Pt(26)

            g_box = guidance_slide.shapes.add_textbox(Inches(0.8), Inches(1.4), Inches(11.5), Inches(5.5))
            gtf = g_box.text_frame
            gtf.word_wrap = True
            
            matrix = implementation_guidance.get("tech_stack_matrix", [])
            for idx, item in enumerate(matrix[:4]):  # Limit to 4 to fit on one slide
                p_item = gtf.paragraphs[0] if idx == 0 else gtf.add_paragraph()
                p_item.text = f"Component: {item.get('component_name')} ({item.get('component_role')})"
                p_item.font.bold = True
                p_item.font.size = Pt(14)
                
                p_tech = gtf.add_paragraph()
                libs = ", ".join(item.get('recommended_libraries_frameworks', []))
                p_tech.text = (
                    f"  - Primary Choice: {item.get('primary_best_tech')} [{item.get('tech_category')}]\n"
                    f"  - Free / Open-Source Alternative: {item.get('open_source_alternative')}\n"
                    f"  - Libraries/SDKs: {libs}"
                )
                p_tech.font.size = Pt(12)

        # Slides 4+: Embed all available diagram types
        slide_diagram_configs = [
            ("architecture_diagram_path", "Component Architecture View"),
            ("usecase_diagram_path", "Functional Scope & Use Case Diagram"),
            ("dfd_diagram_path", "Data Flow & Processing Pipeline"),
            ("er_diagram_path", "Entity-Relationship & Data Model"),
            ("deployment_diagram_path", "Cloud Infrastructure Topography"),
            ("sequence_diagram_path", "Critical Sequence Interactions"),
            ("state_diagram_path", "Entity Lifecycle State Machine"),
        ]

        for key, slide_title in slide_diagram_configs:
            img_path = diagrams.get(key)
            if img_path and os.path.exists(img_path):
                diag_slide = prs.slides.add_slide(blank_layout)
                d_hdr = diag_slide.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(11.0), Inches(0.8))
                dh_p = d_hdr.text_frame.paragraphs[0]
                dh_p.text = slide_title
                dh_p.font.bold = True
                dh_p.font.size = Pt(26)
                
                diag_slide.shapes.add_picture(
                    img_path, 
                    left=Inches(1.5), 
                    top=Inches(1.5), 
                    width=Inches(10.33)
                )

        output_pptx_path = os.path.join(self.output_dir, filename)
        prs.save(output_pptx_path)
        return output_pptx_path

    # =====================================================================
    # Word (.docx) Compiler
    # =====================================================================
    def create_docx_report(
        self, 
        hld_data: Dict[str, Any], 
        diagram_paths: Optional[Dict[str, Optional[str]]] = None,
        implementation_guidance: Optional[Dict[str, Any]] = None,
        filename: str = "HLD_Specification.docx"
    ) -> str:
        """Generates a structured Word document report from HLD schema data, implementation guidance, and diagrams."""
        diagrams = diagram_paths or {}
        doc = Document()
        
        # Title
        doc.add_heading(hld_data.get("system_name", "High-Level Architecture Specification"), level=0)
        
        # 1. Summary
        doc.add_heading("1. Executive Summary", level=1)
        doc.add_paragraph(hld_data.get("summary", ""))
        doc.add_paragraph(f"Architectural Pattern: {hld_data.get('architectural_pattern', 'N/A')}")

        # 2. Components Table
        doc.add_heading("2. Component Specification", level=1)
        components = hld_data.get("components", [])
        if components:
            table = doc.add_table(rows=1, cols=3)
            table.style = 'Table Grid'
            hdr_cells = table.rows[0].cells
            hdr_cells[0].text = 'Component Name'
            hdr_cells[1].text = 'Technology Stack'
            hdr_cells[2].text = 'Description'
            
            for comp in components:
                row_cells = table.add_row().cells
                row_cells[0].text = str(comp.get('name', ''))
                row_cells[1].text = str(comp.get('tech_stack', ''))
                row_cells[2].text = str(comp.get('description', ''))

        # 3. Detailed Implementation Architecture Guidance
        if implementation_guidance:
            doc.add_heading("3. Detailed Implementation Tech Stack & OSS Alternatives", level=1)
            
            if implementation_guidance.get("architecture_summary"):
                doc.add_paragraph(implementation_guidance.get("architecture_summary"))

            matrix = implementation_guidance.get("tech_stack_matrix", [])
            
            if matrix:
                table = doc.add_table(rows=1, cols=4)
                table.style = 'Table Grid'
                hdr = table.rows[0].cells
                hdr[0].text = 'Component'
                hdr[1].text = 'Primary Tech (Best Choice)'
                hdr[2].text = 'Free / Open-Source Alternative'
                hdr[3].text = 'Recommended Libraries'
                
                for item in matrix:
                    row = table.add_row().cells
                    comp_name = item.get('component_name', '')
                    comp_role = item.get('component_role', '')
                    row[0].text = f"{comp_name}\n({comp_role})"
                    
                    primary = item.get('primary_best_tech', '')
                    rationale = item.get('selection_rationale', '')
                    row[1].text = f"{primary}\n\nRationale: {rationale}"
                    
                    alt = item.get('open_source_alternative', '')
                    alt_rat = item.get('alternative_rationale', '')
                    row[2].text = f"{alt}\n\nTrade-offs: {alt_rat}"
                    
                    libs = item.get('recommended_libraries_frameworks', [])
                    if isinstance(libs, list):
                        row[3].text = ", ".join(libs)
                    else:
                        row[3].text = str(libs)

            roadmap = implementation_guidance.get("implementation_roadmap_notes", [])
            if roadmap:
                doc.add_heading("Implementation Notes & Roadmap", level=2)
                for note in roadmap:
                    doc.add_paragraph(str(note), style='List Bullet')

        # 4. All 7 Diagram Sections
        doc_diagram_configs = [
            ("architecture_diagram_path", "4. Component Architecture View"),
            ("usecase_diagram_path", "5. Functional Scope & Use Case Diagram"),
            ("dfd_diagram_path", "6. Data Flow & Processing Pipeline"),
            ("er_diagram_path", "7. Entity-Relationship & Data Model"),
            ("deployment_diagram_path", "8. Cloud Infrastructure Topography"),
            ("sequence_diagram_path", "9. Critical Sequence Interactions"),
            ("state_diagram_path", "10. Entity Lifecycle State Machine"),
        ]

        for key, section_title in doc_diagram_configs:
            img_path = diagrams.get(key)
            if img_path and os.path.exists(img_path):
                doc.add_heading(section_title, level=1)
                doc.add_picture(img_path, width=DocxInches(6.0))

        # 5. Infrastructure & Security Specs
        infra = hld_data.get("infrastructure", {})
        if infra:
            doc.add_heading("11. Infrastructure & Security Specs", level=1)
            doc.add_paragraph(f"Cloud Provider: {infra.get('cloud_provider', 'N/A')}")
            doc.add_paragraph(f"Scaling Strategy: {infra.get('scaling_strategy', 'N/A')}")
            
            core_services = infra.get("core_services", [])
            if core_services:
                doc.add_paragraph(f"Core Services: {', '.join(core_services)}")
                
            sec_controls = infra.get("security_controls", [])
            if sec_controls:
                doc.add_paragraph(f"Security Controls: {', '.join(sec_controls)}")

        output_docx_path = os.path.join(self.output_dir, filename)
        doc.save(output_docx_path)
        return output_docx_path

# =====================================================================
# Standalone Execution / Test Routine
# =====================================================================
if __name__ == "__main__":
    compiler = DocumentCompilerTool()
    
    mock_hld_dict = {
        "system_name": "E-Commerce Vector Search API",
        "architectural_pattern": "Event-Driven Microservices",
        "summary": "High-throughput semantic vector search infrastructure running on AWS ECS and OpenSearch.",
        "components": [
            {
                "name": "API Gateway",
                "tech_stack": "AWS CloudFront + ALB",
                "description": "Routes external REST traffic and validates authorization JWTs."
            }
        ],
        "infrastructure": {
            "cloud_provider": "AWS",
            "scaling_strategy": "Auto-scaling ECS Fargate task instances."
        }
    }
    
    mock_guidance_dict = {
        "system_title": "Implementation Guidance - Vector Search",
        "architecture_summary": "Implementation tech stack tagging with open-source options for self-hosting.",
        "tech_stack_matrix": [
            {
                "component_name": "Vector Database",
                "component_role": "Embeddings storage and sub-50ms approximate nearest neighbor search",
                "primary_best_tech": "Pinecone (Serverless)",
                "tech_category": "Managed SaaS",
                "is_proprietary_or_paid": True,
                "selection_rationale": "Zero-ops scaling with optimized dense-sparse hybrid indexing.",
                "open_source_alternative": "Qdrant / Milvus",
                "alternative_rationale": "Can be self-hosted on Kubernetes clusters with minimal memory overhead.",
                "recommended_libraries_frameworks": ["qdrant-client", "fastembed", "langchain-community"]
            }
        ],
        "implementation_roadmap_notes": [
            "Set up Qdrant Docker container or Pinecone API keys first.",
            "Implement ingestion pipeline worker prior to exposing search API."
        ]
    }
    
    print("Testing DocumentCompilerTool locally...")
    try:
        pptx_path = compiler.create_pptx_deck(
            mock_hld_dict, 
            implementation_guidance=mock_guidance_dict, 
            filename="Test_HLD.pptx"
        )
        docx_path = compiler.create_docx_report(
            mock_hld_dict, 
            implementation_guidance=mock_guidance_dict, 
            filename="Test_HLD.docx"
        )
        print("\n--- Deliverables Successfully Compiled! ---")
        print(f"PowerPoint Deck: {os.path.abspath(pptx_path)}")
        print(f"Word Specification: {os.path.abspath(docx_path)}")
    except Exception as err:
        print(f"Compilation error: {err}")